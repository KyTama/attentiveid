import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json, os, csv, re

REGISTRY_PATH = '.docs/domain/psychologists-registry.json'
with open(REGISTRY_PATH, 'r', encoding='utf-8') as f:
    registry = json.load(f)

# Create Workbook
wb = openpyxl.Workbook()

# Styles
font_title = Font(name='Calibri', size=14, bold=True, color='142C52')
font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
font_bold = Font(name='Calibri', size=10, bold=True)
font_regular = Font(name='Calibri', size=10)
font_italic = Font(name='Calibri', size=9, italic=True, color='555555')

fill_navy = PatternFill(start_color='142C52', end_color='142C52', fill_type='solid')
fill_gold = PatternFill(start_color='D6A74D', end_color='D6A74D', fill_type='solid')
fill_light_cream = PatternFill(start_color='FEFAF6', end_color='FEFAF6', fill_type='solid')
fill_light_gray = PatternFill(start_color='F5F5F7', end_color='F5F5F7', fill_type='solid')
fill_alert_red = PatternFill(start_color='FFEBEE', end_color='FFEBEE', fill_type='solid')
fill_alert_green = PatternFill(start_color='E8F5E9', end_color='E8F5E9', fill_type='solid')
fill_alert_yellow = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')

border_thin = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)

align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

# -------------------------------------------------------------
# SHEET 1: 01_Data_Utama_Psikolog
# -------------------------------------------------------------
ws1 = wb.active
ws1.title = '01_Data_Utama_Psikolog'
ws1.views.sheetView[0].showGridLines = True

headers1 = [
    'No', 'Slug ID', 'Nama Lengkap & Gelar', 'Panggilan', 'Gender',
    'Status Praktik', 'Status Kontrak', 'Tiering', 'Menerima Klien Baru',
    'Tahun Pengalaman', 'Tahun Gabung', 'Lokasi Cabang Utama', 'Format Layanan',
    'Peminatan Utama', 'Domain Layanan (Support Areas)', 'Populasi Klien / Usia',
    'No. Izin Praktik Terkini (SIPP / STR / SIPPK)',
    'Pendidikan (Kampus)', 'Tempat Kerja Lain', 'Status Foto Studio',
    'Dokumen CV / Portofolio', 'Catatan Data Kurang / To-Do Tim Ops'
]

# Title row
ws1.append(['DATABASE & AUDIT PSIKOLOG BIRO ATTENTIVE 2026'])
ws1.merge_cells('A1:V1')
ws1['A1'].font = font_title
ws1['A1'].alignment = Alignment(horizontal='left', vertical='center')
ws1.row_dimensions[1].height = 30

ws1.append(['Gunakan sheet ini untuk konfirmasi status aktif, kelengkapan nomor izin terkini, serta foto portrait studio sebelum verifikasi database.'])
ws1.merge_cells('A2:V2')
ws1['A2'].font = font_italic
ws1.row_dimensions[2].height = 20

# Header row
ws1.append(headers1)
header_row_idx = 3
ws1.row_dimensions[header_row_idx].height = 28

for col_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=header_row_idx, column=col_idx)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center
    cell.border = border_thin

# Specific override logic for status aktif vs nonaktif
def get_operational_status(p):
    slug = p['id']
    if slug == 'kia':
        return 'Nonaktif (Resign / Tidak Praktik)'
    if slug == 'gisella':
        return 'Pending Review (Afiliasi Psycoach / Izin Proses)'
    return 'Aktif'

def get_contract_status(p):
    slug = p['id']
    if slug == 'kia':
        return 'Nonaktif'
    if slug == 'gisella':
        return 'Belum Ada Kontrak'
    if slug == 'haykal':
        return 'Partner / Consultant'
    return 'Sudah TTD'

def get_missing_notes(p):
    missing = []
    slug = p['id']
    license_num = p.get('licenseNumber') or ''
    if slug == 'kia':
        return 'Status kontrak Nonaktif (Resign); data dipertahankan di DB (status inactive, accepting_new_clients false) demi integritas relasi sesi lama.'
    if slug == 'gisella':
        missing.append('Konfirmasi kontrak kemitraan Attentive & nomor SIPPK (catatan: sedang diurus)')
    if not p['assets']['hasPhoto'] or 'NEED_PHOTO' in p['readinessStatus']:
        missing.append('Foto portrait studio resolusi tinggi belum ada di folder Photos/')
    if not license_num or license_num in ['belum ada', 'Harus Update', 'Tidak ada', None, '-', '(kosong)'] or 'harus update' in license_num.lower():
        missing.append(f"Nomor izin resmi terbaru ({license_num or 'kosong'})")
    if not p['education']['campus'] or p['education']['campus'] in ['-', '(kosong)']:
        missing.append('Detail nama kampus S1 / S2')
    if not p['landingDetails'].get('shortFitStatement'):
        missing.append('Short fit statement 1-2 kalimat untuk kartu landing page')
    if not p['landingDetails'].get('primaryApproaches'):
        missing.append('Daftar pendekatan klinis yang aktif digunakan')
        
    return '; '.join(missing) if missing else 'Lengkap & Siap Live ✅'

tier_labels = {
    'principal': 'Principal',
    'senior': 'Senior',
    'senior_mid': 'Senior-Mid',
    'mid': 'Mid',
    'consultant': 'Consultant'
}

for idx, p in enumerate(registry, 1):
    op_status = get_operational_status(p)
    contract_st = get_contract_status(p)
    missing_notes = get_missing_notes(p)
    
    photo_text = 'Ada (Photos/)' if p['assets']['rawPhotoPath'] else 'Ada (Web Asset)' if p['assets']['hasPhoto'] else 'Belum Ada ❌'
    cv_text = 'Ada (PDF) ✅' if p['assets']['hasCv'] else 'Belum Ada'
    support_areas_str = ', '.join(p['supportAreas'])
    pop_str = p['landingDetails'].get('targetPopulation') or ('Dewasa 18+' if p['peminatan'] == 'Klinis Dewasa' else 'Anak & Remaja' if 'Anak' in p['peminatan'] else 'Pelajar & Mahasiswa' if 'Pendidikan' in p['peminatan'] else 'Umum')
    
    branch_display = p.get('primaryLocationDisplay') or p.get('primaryBranch', 'TBI')
    
    fmt_layanan = 'Online & Offline'
    if p.get('primaryBranch') == 'online_only':
        fmt_layanan = 'Online Only'
    elif p.get('primaryBranch') == 'malang':
        fmt_layanan = 'Offline (Malang) & Online'
    elif p.get('primaryBranch') in ['tbi', 'bsd']:
        fmt_layanan = f'Offline ({branch_display}) & Online'
    elif p.get('primaryBranch') == 'multiple':
        fmt_layanan = 'Offline (TBI & BSD) & Online'
        
    accepting_clients_str = 'Ya ✅' if p.get('acceptingNewClients', True) else 'Tidak (Ditutup) ⛔'
    tier_display = tier_labels.get(p.get('tier', 'mid'), p.get('tier', 'mid').title())
    
    row_data = [
        idx,
        p['id'],
        p['fullName'],
        p['nickname'],
        'Perempuan' if p['gender'] == 'female' else 'Laki-laki',
        op_status,
        contract_st,
        tier_display,
        accepting_clients_str,
        p['experienceYears'],
        p['joinDate'],
        branch_display,
        fmt_layanan,
        p['peminatan'],
        support_areas_str,
        pop_str,
        p.get('licenseNumber') or '-',
        p['education']['campus'] or '-',
        p['workplace'] or '-',
        photo_text,
        cv_text,
        missing_notes
    ]
    ws1.append(row_data)
    row_num = ws1.max_row
    ws1.row_dimensions[row_num].height = 24
    
    # Cell formatting
    for col_idx in range(1, len(row_data) + 1):
        c = ws1.cell(row=row_num, column=col_idx)
        c.font = font_regular
        c.border = border_thin
        if col_idx in [1, 2, 4, 5, 8, 9, 10, 11, 12, 20, 21]:
            c.alignment = align_center
        else:
            c.alignment = align_left
            
        # Highlight status aktif
        if col_idx == 6:
            if 'Aktif' in op_status and 'Nonaktif' not in op_status:
                c.fill = fill_alert_green
                c.font = font_bold
            elif 'Nonaktif' in op_status:
                c.fill = fill_alert_red
                c.font = font_bold
            else:
                c.fill = fill_alert_yellow
                c.font = font_bold
                
        # Highlight accepting clients
        if col_idx == 9:
            if 'Ya' in accepting_clients_str:
                c.fill = fill_alert_green
            else:
                c.fill = fill_alert_red
                
        # Highlight photo
        if col_idx == 20:
            if 'Belum' in photo_text:
                c.fill = fill_alert_yellow
                
        # Highlight notes
        if col_idx == 22:
            if 'Lengkap' in missing_notes:
                c.fill = fill_alert_green
            else:
                c.fill = fill_alert_yellow

# -------------------------------------------------------------
# SHEET 2: 02_Expertise_&_Problem_List
# -------------------------------------------------------------
ws2 = wb.create_sheet(title='02_Expertise_&_Problem_List')
ws2.views.sheetView[0].showGridLines = True

headers2 = [
    'No', 'Slug ID', 'Nama Lengkap', 'Panggilan', 'Peminatan', 'Tiering', 'Cabang Utama',
    'Topik 1', 'Topik 2', 'Topik 3', 'Topik 4', 'Topik 5', 'Topik 6', 'Topik 7', 'Topik 8',
    'Pendekatan Terapi / Modalitas', 'Short Fit Statement (Untuk Kartu Profil)', 'Gaya Interaksi / Working Style'
]

ws2.append(['MAPPING KELEBIHAN KLINIS & TOPIK MASALAH (EXPERTISE MATRIX)'])
ws2.merge_cells('A1:R1')
ws2['A1'].font = font_title
ws2.row_dimensions[1].height = 30

ws2.append(['Data topik masalah dan pendekatan klinis ini digunakan oleh Algoritma Matching Intake Survey & filter direktori publik psikolog.'])
ws2.merge_cells('A2:R2')
ws2['A2'].font = font_italic
ws2.row_dimensions[2].height = 20

ws2.append(headers2)
ws2.row_dimensions[3].height = 28

for col_idx, h in enumerate(headers2, 1):
    cell = ws2.cell(row=3, column=col_idx)
    cell.font = font_header
    cell.fill = fill_gold
    cell.alignment = align_center
    cell.border = border_thin

for idx, p in enumerate(registry, 1):
    topics = p['expertiseTopics'] + [''] * 8
    topics_8 = topics[:8]
    branch_display = p.get('primaryLocationDisplay') or p.get('primaryBranch', 'TBI')
    tier_display = tier_labels.get(p.get('tier', 'mid'), p.get('tier', 'mid').title())
    
    row_data = [
        idx,
        p['id'],
        p['fullName'],
        p['nickname'],
        p['peminatan'],
        tier_display,
        branch_display,
        topics_8[0], topics_8[1], topics_8[2], topics_8[3],
        topics_8[4], topics_8[5], topics_8[6], topics_8[7],
        p['landingDetails'].get('primaryApproaches') or '-',
        p['landingDetails'].get('shortFitStatement') or '-',
        p['landingDetails'].get('workingStyle') or '-'
    ]
    ws2.append(row_data)
    row_num = ws2.max_row
    ws2.row_dimensions[row_num].height = 22
    for col_idx in range(1, len(row_data) + 1):
        c = ws2.cell(row=row_num, column=col_idx)
        c.font = font_regular
        c.border = border_thin
        if col_idx in [1, 2, 4, 6, 7]:
            c.alignment = align_center
        else:
            c.alignment = align_left

# -------------------------------------------------------------
# SHEET 3: 03_Audit_Database_Schema_Gap
# -------------------------------------------------------------
ws3 = wb.create_sheet(title='03_Audit_Schema_Gap')
ws3.views.sheetView[0].showGridLines = True

headers3 = ['No', 'Komponen / Field Domain', 'Status di Database PostgreSQL', 'Kebutuhan Nyata di Lapangan & Keputusan User', 'Status Implementasi', 'Implementasi Teknis Drizzle / Postgres']

ws3.append(['STATUS IMPLEMENTASI & AUDIT SKEMA DATABASE (POSTGRES / DRIZZLE)'])
ws3.merge_cells('A1:F1')
ws3['A1'].font = font_title
ws3.row_dimensions[1].height = 30

ws3.append(['Laporan status keselarasan skema database (apps/api/src/db/schema.ts) dengan keputusan operasional & hasil migrasi 0004.'])
ws3.merge_cells('A2:F2')
ws3['A2'].font = font_italic
ws3.row_dimensions[2].height = 20

ws3.append(headers3)
ws3.row_dimensions[3].height = 28

for col_idx, h in enumerate(headers3, 1):
    cell = ws3.cell(row=3, column=col_idx)
    cell.font = font_header
    cell.fill = fill_navy
    cell.alignment = align_center
    cell.border = border_thin

gap_items = [
    (
        1,
        'Nomor Izin Praktik (SIPP / STR / SIPPK)',
        'Kolom tunggal text: license_number pada psychologist_profiles',
        'Keputusan User: Tetap satukan dalam satu kolom text. Jika memiliki multi-regulasi (SIPP, STR, SIPPK), gunakan izin paling terkini dan aktif.',
        'TERSELESAIKAN (Done)',
        'Kolom license_number diisi lisensi terkini (misal SIPPK 2025 s.d. 2030 untuk Jeanette, STR-PK untuk Nuzul, SIPP 2025 untuk Andri).'
    ),
    (
        2,
        'Lokasi Cabang & Praktik (TBI / BSD / Malang / Online)',
        'Kolom primary_branch enum ("tbi", "bsd", "malang", "online_only", "multiple") di tabel psychologists',
        'Attentive memiliki 3 cabang fisik (Tanjung Barat, BSD, Malang), serta sesi Online. Klien butuh filter cabang saat reservasi.',
        'TERSELESAIKAN (Done)',
        'Ditambahkan practice_branch_enum dan kolom primary_branch (default "tbi") di migrasi 0004. Index ditambahkan untuk optimasi query cabang.'
    ),
    (
        3,
        'Tiering Psikolog (Principal s.d. Mid)',
        'Kolom tier enum ("principal", "senior", "senior_mid", "mid", "consultant") di tabel psychologists',
        'Tarif konsultasi dan level kompetensi ditentukan oleh tiering psikolog. Klien melihat tarif berbeda berdasarkan tier.',
        'TERSELESAIKAN (Done)',
        'Ditambahkan psychologist_tier_enum dan kolom tier (default "mid") di migrasi 0004. Digunakan untuk perhitungan tarif konsultasi.'
    ),
    (
        4,
        'Status Resign & Integritas Relasi (Inactivating vs Deletion)',
        'Kolom status enum ("draft", "active", "inactive", "archived") dan accepting_new_clients (boolean)',
        'Psikolog yang resign (seperti Riskia Murad) tidak boleh di-DELETE dari database agar riwayat konsultasi klien tetap memiliki integritas referensial (FK RESTRICT).',
        'TERSELESAIKAN (Done)',
        'Ditambahkan kolom accepting_new_clients (boolean default true). Psikolog nonaktif ditandai status="inactive" dan accepting_new_clients=false.'
    ),
    (
        5,
        'Pemetaan Topik Masalah & Intake Matching Engine',
        'Tag array di specializations dan rules di intake matching engine',
        'Intake survey menyaring keluhan pasien (gangguan mood, trauma, relasi, dll.) dan memetakan ke psikolog yang menguasai topik tersebut.',
        'TERPETAKAN (Ready to Seed)',
        'Data 24 psikolog memiliki 8 topik keahlian utama di registry JSON. Siap disinkronkan ke tabel database & filter intake.'
    ),
    (
        6,
        'Pendekatan Terapi (Therapeutic Approaches) & Short Fit Statement',
        'Field biography dan translation notes di database',
        'Kartu profil psikolog di web membutuhkan short fit statement ("Mungkin relevan jika...") dan daftar modalitas (CBT, SE, Art Therapy).',
        'TERPETAKAN (Ready to Seed)',
        'Field shortFitStatement dan primaryApproaches telah terstruktur di registry JSON dan data center markdown.'
    )
]

for item in gap_items:
    ws3.append(list(item))
    row_num = ws3.max_row
    ws3.row_dimensions[row_num].height = 36
    for col_idx in range(1, len(item) + 1):
        c = ws3.cell(row=row_num, column=col_idx)
        c.font = font_regular
        c.border = border_thin
        if col_idx in [1, 5]:
            c.alignment = align_center
            if 'Done' in item[4]:
                c.fill = fill_alert_green
                c.font = font_bold
            else:
                c.fill = fill_alert_yellow
                c.font = font_bold
        else:
            c.alignment = align_left

# Auto-adjust column widths across all sheets
for sheet in [ws1, ws2, ws3]:
    for col in sheet.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

EXCEL_OUT = '.docs/domain/Attentive_Psychologist_Master_Database_Audit_2026.xlsx'
wb.save(EXCEL_OUT)
print(f"Successfully generated Excel: {EXCEL_OUT}")

# Also export Sheet 1 to CSV for instant 1-click Google Sheets import
CSV_OUT = '.docs/domain/Attentive_Psychologist_Master_Database_Audit_2026.csv'
with open(CSV_OUT, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    for row in ws1.iter_rows(values_only=True):
        if row and any(row):
            writer.writerow(row)
print(f"Successfully generated CSV: {CSV_OUT}")
